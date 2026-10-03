import os

import pandas as pd
from openpyxl import load_workbook

from reporte import generar_reporte


def test_genera_archivo(tmp_path):
    csv = tmp_path / "in.csv"
    csv.write_text("a,b\n1,10\n2,20\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    generar_reporte(str(csv), str(out))
    assert out.exists() and out.stat().st_size > 0


def test_hojas_correctas(tmp_path):
    csv = tmp_path / "in.csv"
    csv.write_text("a,b\n1,10\n2,20\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    generar_reporte(str(csv), str(out))
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"Datos", "Resumen"}


def test_totales_correctos(tmp_path):
    csv = tmp_path / "in.csv"
    csv.write_text("a,b\n1,10\n2,20\n3,30\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    stats = generar_reporte(str(csv), str(out))
    assert stats["filas"] == 3
    resumen = pd.read_excel(out, sheet_name="Resumen", index_col=0)
    assert resumen.loc["a", "total"] == 6
    assert resumen.loc["b", "promedio"] == 20
