"""Genera un reporte Excel (datos + resumen) a partir de un CSV.

Uso: python reporte.py entrada.csv --salida reporte.xlsx
"""
import argparse

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def generar_reporte(csv_path: str, salida: str) -> dict:
    """Lee el CSV y produce un xlsx con hojas 'Datos' y 'Resumen'.

    Devuelve un dict con métricas del resumen (útil para tests/logging).
    """
    df = pd.read_csv(csv_path)
    num = df.select_dtypes(include="number")

    resumen = pd.DataFrame({
        "total": num.sum(),
        "promedio": num.mean(),
        "máximo": num.max(),
        "mínimo": num.min(),
    })

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Datos", index=False)
        resumen.to_excel(writer, sheet_name="Resumen")

        wb = writer.book
        for sheet in ("Datos", "Resumen"):
            ws = wb[sheet]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col)].width = 18

    return {"filas": len(df), "columnas_numericas": list(num.columns)}


def main():
    p = argparse.ArgumentParser(description="CSV → reporte Excel con resumen")
    p.add_argument("csv", help="Archivo CSV de entrada")
    p.add_argument("--salida", default="reporte.xlsx", help="XLSX de salida")
    args = p.parse_args()

    stats = generar_reporte(args.csv, args.salida)
    print(f"OK: {args.salida} generado "
          f"({stats['filas']} filas, columnas numéricas: {stats['columnas_numericas']})")


if __name__ == "__main__":
    main()
